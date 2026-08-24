"""One standalone, editable slide: DRISHTI's technical approach and data flow.

    python scripts/build_technical_approach_slide.py

Deliberately **not** part of ``build_deck.py`` and does not import it: that
script executes its whole six-slide build at import time and saves over the
existing submission deck, which is frequently open in PowerPoint and may hold
edits made by hand since the last regeneration. Overwriting it as a side
effect of building one extra slide would be exactly the kind of destructive
surprise this project avoids elsewhere. This script owns its own small copy of
the same drawing primitives instead, and writes to its own file.

Layout mirrors the reference the request was built from: a title bar, a
bordered "how it's built" text panel on the left, and a labelled process-flow
diagram on the right. The diagram is not decorative — every box and arrow
matches a real module in this codebase, named beside it, so nothing on the
slide is a claim the code cannot back up.

Open the output in PowerPoint or Google Slides; every shape and text run is a
native editable object, not a picture.
"""

from __future__ import annotations

import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as I
from pptx.util import Pt

# Real font-metric measurement, copied from build_deck.py's own solution to
# the same problem: a guessed line height is how that deck shipped three
# silent overflows the first time it was built. Duplicated rather than
# imported - see the module docstring for why importing that script is unsafe.
_FONTFILE = {("Calibri", False): r"C:\Windows\Fonts\calibri.ttf",
             ("Calibri", True):  r"C:\Windows\Fonts\calibrib.ttf",
             ("Cambria", False): r"C:\Windows\Fonts\cambria.ttc",
             ("Cambria", True):  r"C:\Windows\Fonts\cambriab.ttf"}
_fcache = {}


def _font(name, bold, pt):
    k = (name, bold, round(pt))
    if k not in _fcache:
        _fcache[k] = ImageFont.truetype(_FONTFILE[(name, bold)], int(round(pt * 4)))
    return _fcache[k]


def nlines(text, width_in, pt, name="Calibri", bold=False):
    f = _font(name, bold, pt)
    limit = width_in * 72.0
    total = 0
    for hard in text.split("\n"):
        line, n = "", 1
        for w in hard.split(" "):
            trial = w if not line else line + " " + w
            if f.getlength(trial) / 4.0 <= limit or not line:
                line = trial
            else:
                n += 1
                line = w
        total += n
    return total


def height_in(text, width_in, pt, name="Calibri", bold=False, spacing=1.16):
    return nlines(text, width_in, pt, name, bold) * pt * 1.20 * spacing / 72.0

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.environ.get("DRISHTI_SLIDE_OUT") or os.path.join(
    ROOT, "docs", "DRISHTI_Technical_Approach_Slide.pptx")

# Same palette and fonts as the main deck (scripts/build_deck.py), copied
# rather than imported so this script has no side effect on that file.
NAVY, SLATE, STEEL = "0B1B2B", "17324A", "24485F"
ICE, WHITE = "EDF2F6", "FFFFFF"
AMBER, RED, GREEN, YELLOW = "E07B39", "C4362C", "2E8B57", "E3B23C"
BLUE = "2E7DB5"
MUTED_D, MUTED_L = "9FB4C7", "5A7186"
HEAD, BODY, MONO = "Cambria", "Calibri", "Consolas"

W, H, M = 13.333, 7.5, 0.5
RIGHT = W - M

prs = Presentation()
prs.slide_width, prs.slide_height = I(W), I(H)
BLANK = prs.slide_layouts[6]
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = C.from_string(WHITE)


def rect(l, t, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, lw=1.0,
        radius=None):
    sh = s.shapes.add_shape(shape, I(l), I(t), I(w), I(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = C.from_string(fill)
    if line:
        sh.line.color.rgb = C.from_string(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def txt(l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = s.shapes.add_textbox(I(l), I(t), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (text, size, bold, col, font, sp) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sp)
        p.line_spacing = 1.12
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = C.from_string(col)
        r.font.name = font
    return tb


def box_label(l, t, w, h, text, fill, line, size=10.5, bold=True, col=WHITE,
             align=PP_ALIGN.CENTER):
    b = rect(l, t, w, h, fill, line=line, lw=1.0, radius=0.12)
    b.text_frame.word_wrap = True
    b.text_frame.margin_left = b.text_frame.margin_right = Pt(4)
    b.text_frame.margin_top = b.text_frame.margin_bottom = Pt(2)
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = b.text_frame.paragraphs[0] if i == 0 else b.text_frame.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size if i == 0 else size - 1.5)
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = C.from_string(col)
        r.font.name = BODY
    return b


def arrow(x1, y1, x2, y2, color=STEEL, weight=1.75):
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    conn.line.color.rgb = C.from_string(color)
    conn.line.width = Pt(weight)
    ln = conn.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


# ---------------------------------------------------------------------------
# header
# ---------------------------------------------------------------------------

rect(0, 0, W, 0.08, AMBER, shape=MSO_SHAPE.RECTANGLE)
txt(0, 0.24, W, 0.6,
    [("TECHNICAL APPROACH", 30, True, NAVY, HEAD, 0)],
    align=PP_ALIGN.CENTER)
txt(0, 0.78, W, 0.34,
    [("DRISHTI — District Risk Intelligence and Satellite Hazard Tracking Interface  ·  SIH26191",
      12, False, MUTED_L, BODY, 0)],
    align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# left panel — system architecture & workflow (text)
# ---------------------------------------------------------------------------

LX, LY, LW, LH = M, 1.35, 5.55, 6.0
rect(LX, LY, LW, LH, WHITE, line="2E5C8A", lw=1.75, radius=0.03)
rect(LX, LY, LW, 0.5, "0B1B2B", radius=0.03)
txt(LX + 0.22, LY + 0.09, LW - 0.4, 0.34,
    [("System Architecture & Workflow", 15, True, WHITE, HEAD, 0)])

body_y = LY + 0.68
BODY_W = LW - 0.44
BODY_PT = 10.5
HEAD_PT = 11.5


def block(heading, bullets, color, y):
    """One heading + its bullets, each sized from real measured line-wrap
    height rather than a guessed constant — the exact bug this replaces."""
    hh = height_in(heading, BODY_W, HEAD_PT, BODY, True)
    txt(LX + 0.22, y, BODY_W, hh + 0.05,
        [(heading, HEAD_PT, True, color, BODY, 0)])
    y += hh + 0.09
    for b in bullets:
        bh = height_in(b, BODY_W, BODY_PT, BODY, False)
        txt(LX + 0.22, y, BODY_W, bh + 0.03,
            [(b, BODY_PT, False, "20364A", BODY, 0)])
        y += bh + 0.06
    return y + 0.11


y = body_y
y = block("FRONTEND — vanilla JS + Leaflet, no build step", [
    "•  Live board, click-to-justify red zones, district & national views",
    "•  One state object drives every screen; no framework, no bundler",
], BLUE, y)

y = block("BACKEND — FastAPI + NumPy physics engine", [
    "•  Terrain hydrology: D8 flow routing, HAND, slope, wetness index",
    "•  Hazard models: BIS IS 14496 (landslide), SCS-CN→Manning (flood), "
    "IMD threshold (cloudburst), Bruun (erosion)",
], AMBER, y)

y = block("DATA LAYER — zero API keys, anywhere", [
    "•  Baked real: Copernicus DEM · GHS-POP · GHSL · WorldCover · "
    "OpenStreetMap · geoBoundaries · IBTrACS",
    "•  Live feeds: Open-Meteo — rainfall, CAPE, soil moisture, cloud structure",
], GREEN, y)

y = block("DATA FLOW", [
    "Terrain + satellite data → per-cell hazard physics → Red Zone "
    "recurrence index → relocation site allocation + live monitoring",
], RED, y)

y = block("LIVE / RESPONSE LOOP", [
    "Live weather → cloud-pattern read (no ML) → urgency ranking → "
    "force-refresh proof → district drill-down → click-to-justify",
], "6A3FA0", y)

# ---------------------------------------------------------------------------
# right panel — process flow architecture (diagram)
# ---------------------------------------------------------------------------

RX, RY, RW = LX + LW + 0.3, 1.35, RIGHT - (LX + LW + 0.3)
txt(RX, RY, RW, 0.3, [("PROCESS FLOW ARCHITECTURE", 13, True, "0B5C8C", BODY, 0)])

fx, fy, fw = RX + 0.1, RY + 0.42, RW - 0.2
row_h = 0.5
gap = 0.28

# Row 1 — two inputs side by side
half = (fw - 0.2) / 2
box_label(fx, fy, half, row_h,
          "BAKED REAL DATA\nDEM · Population · Built-up · Roads",
          GREEN, "1E6B44", size=9.5)
box_label(fx + half + 0.2, fy, half, row_h,
          "LIVE FEEDS (Open-Meteo)\nRainfall · CAPE · Soil · Cloud",
          BLUE, "1E5A85", size=9.5)
y1 = fy + row_h

arrow(fx + half / 2, y1, fx + fw / 2 - 0.35, y1 + gap - 0.03)
arrow(fx + half + 0.2 + half / 2, y1, fx + fw / 2 + 0.35, y1 + gap - 0.03)

# Row 2 — terrain engine
y2 = y1 + gap
box_label(fx, y2, fw, row_h * 0.8,
          "TERRAIN ENGINE  —  core/terrain.py\nD8 flow routing · HAND · slope · topographic wetness index",
          STEEL, "0B1B2B", size=9.5)
y3s = y2 + row_h * 0.8

arrow(fx + fw / 2, y3s, fx + fw / 2, y3s + gap - 0.04)

# Row 3 — hazard physics (wider box, 2 lines)
y3 = y3s + gap
hazard_h = row_h * 0.95
box_label(fx, y3, fw, hazard_h,
          "HAZARD PHYSICS  —  five modules, one recurrence scale\n"
          "Flood · Landslide (BIS LHEF) · Cloudburst (IMD) · Erosion · Waterlogging",
          AMBER, "9A4A15", size=9.5)
y4s = y3 + hazard_h

arrow(fx + fw / 2, y4s, fx + fw / 2, y4s + gap - 0.04)

# Row 4 — red zone engine
y4 = y4s + gap
box_label(fx, y4, fw, row_h * 0.8,
          "RED ZONE ENGINE  —  core/redzone.py\nShortest return period (5 / 25 / 100-yr) at which a cell turns uninhabitable",
          RED, "7A2018", size=9.5)
y5s = y4 + row_h * 0.8

# fan out into 3
third = (fw - 0.4) / 3
cx = [fx + third / 2, fx + third + 0.2 + third / 2, fx + 2 * (third + 0.2) + third / 2]
for cxi in cx:
    arrow(fx + fw / 2, y5s, cxi, y5s + gap - 0.03)

# Row 5 — three outputs
y5 = y5s + gap
box_label(fx, y5, third, row_h,
          "RELOCATION ENGINE\nSiting + capacity math",
          "6A3FA0", "3E2266", size=9)
box_label(fx + third + 0.2, y5, third, row_h,
          "LIVE WATCH BOARD\nUrgency ranking, live",
          BLUE, "1E5A85", size=9)
box_label(fx + 2 * (third + 0.2), y5, third, row_h,
          "CLICK-TO-JUSTIFY\nPer-cell explain API",
          GREEN, "1E6B44", size=9)
y6s = y5 + row_h

for cxi in cx:
    arrow(cxi, y6s, fx + fw / 2, y6s + gap - 0.04)

# Row 6 — frontend / decision
y6 = y6s + gap
box_label(fx, y6, fw, row_h * 0.85,
          "FRONTEND — Leaflet map + panels\n→ District Magistrate / SDMA decision",
          NAVY, "000000", size=10)

prs.save(OUT)
print("saved:", OUT)
